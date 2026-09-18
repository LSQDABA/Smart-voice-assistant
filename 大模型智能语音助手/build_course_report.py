from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = r"C:\Users\li\Documents\大模型智能语音助手\嘎嘎鸭智能助手-嵌入式技术课程设计说明书.docx"


def set_run_font(run, font="宋体", size=None, bold=None, color=None):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7C3D0"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.first_line_indent = Pt(22)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color, before, after in [
        ("Heading 1", 15, "1F4E79", 12, 8),
        ("Heading 2", 13, "1F4E79", 10, 6),
        ("Heading 3", 12, "365F91", 8, 4),
    ]:
        st = doc.styles[name]
        st.font.name = "黑体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.first_line_indent = Pt(0)
        st.paragraph_format.line_spacing = 1.3
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)


def add_center_paragraph(doc, text, size=12, bold=False, after=6, font="宋体"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_run_font(r, font=font, size=size, bold=bold)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, size=11)
    return p


def add_noindent(doc, text, bold=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run(text)
    set_run_font(r, size=11, bold=bold, color=color)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p


def add_code(doc, code):
    for line in code.strip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.line_spacing = 1.1
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(line)
        set_run_font(r, font="Consolas", size=9.5)


def add_placeholder(doc, caption, note="请在此处插入对应截图或照片"):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    set_table_borders(table, "9EB6CE")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F7FB")
    set_cell_width(cell, 15.5)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run(f"{caption}\n【图片占位】{note}")
    set_run_font(r, size=11, bold=True, color="4A5A6A")
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Pt(0)
    cap.paragraph_format.space_after = Pt(8)
    rr = cap.add_run(caption)
    set_run_font(rr, size=10, bold=True)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_table_borders(table)
    for idx, h in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "E8EEF5")
        set_cell_width(cell, widths[idx])
        set_cell_text(cell, h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        for idx, val in enumerate(row):
            set_cell_width(cells[idx], widths[idx])
            set_cell_text(cells[idx], val, align=WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph()
    return table


def add_cover(doc):
    add_center_paragraph(doc, "广州应用科技学院", 18, True, 14, "黑体")
    add_center_paragraph(doc, "嵌入式技术课程设计说明书", 20, True, 30, "黑体")
    add_center_paragraph(doc, "嵌入式大模型应用与开发", 17, True, 20, "黑体")
    add_center_paragraph(doc, "课设题目：嘎嘎鸭智能助手", 16, True, 36, "黑体")

    rows = [
        ("学院", "人工智能与电气工程学院（智能制造学院）"),
        ("专业", "电子信息工程"),
        ("班级", "________________"),
        ("学号", "________________"),
        ("学生姓名", "________________"),
        ("指导教师", "________________"),
        ("提交日期", "2026年6月____日"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = False
    for i, (k, v) in enumerate(rows):
        set_cell_width(table.cell(i, 0), 4.0)
        set_cell_width(table.cell(i, 1), 10.0)
        set_cell_text(table.cell(i, 0), k, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.cell(i, 1), v, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_table_borders(table, "FFFFFF")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_toc(doc):
    add_center_paragraph(doc, "目录", 16, True, 14, "黑体")
    entries = [
        "1 系统设计与环境环境安装",
        "1.1 设计任务及要求",
        "1.2 整体设计方案",
        "1.3 GEC6818平台配置",
        "2 Linux与嵌入式开发板使用",
        "2.1 Linux交叉开发环境配置",
        "2.2 SSH工具配置",
        "3 屏幕软件设计",
        "3.1 主程序设计流程",
        "3.2 UI界面交互程序设计",
        "3.3 数据显示程序设计",
        "4 大模型部署与应用",
        "4.1 大模型部署",
        "4.2 大模型应用",
        "4.3 百度语音识别模型开发",
        "5 成果验收与总结",
        "5.1 成果验收",
        "5.2 设计总结",
        "参考资料",
    ]
    for e in entries:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Cm(0.6 if "." in e[:4] else 0)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(e)
        set_run_font(r, size=11)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_header_footer(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.first_line_indent = Pt(0)
    r = header.add_run("嵌入式大模型应用与开发课程设计说明书")
    set_run_font(r, size=9, color="666666")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    rr = footer.add_run("—  ")
    set_run_font(rr, size=9, color="666666")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    rr._r.append(fld_begin)
    rr._r.append(instr)
    rr._r.append(fld_end)
    rrr = footer.add_run("  —")
    set_run_font(rrr, size=9, color="666666")


def build():
    doc = Document()
    style_doc(doc)
    add_header_footer(doc)
    add_cover(doc)
    add_toc(doc)

    doc.add_heading("1 系统设计与环境环境安装", level=1)
    doc.add_heading("1.1 设计任务及要求", level=2)
    add_body(doc, "本课程设计围绕“嵌入式大模型应用与开发”主题，完成一个可运行在GEC6818开发板上的中文智能语音助手系统。系统命名为“嘎嘎鸭智能助手”，要求在嵌入式Linux环境下实现图形界面显示、语音录入、语音识别、文本问题提交以及大模型回答展示等功能。用户通过开发板触摸屏点击语音输入按钮后，系统调用录音工具采集语音，将音频整理为百度语音识别接口需要的WAV格式，再把识别出的文字显示在问题输入框中；用户确认后点击发送按钮，程序把问题提交到扣子智能体服务，最后将大模型返回的回答显示在界面回答区域。")
    add_body(doc, "从功能要求看，本设计需要满足四个方面：第一，界面布局应符合题目给出的示意图，包含标题、模型回答显示区、语音输入按钮、问题输入框和发送按钮；第二，语音识别流程应能在开发板上真实执行，并能处理开发板声卡只支持48kHz双声道录音等实际问题；第三，大模型调用应能通过网络接口完成，并对流式返回结果进行解析；第四，系统应具备一定的异常提示能力，当麦克风、网络、令牌或接口返回异常时，能够在屏幕上给出可读提示，方便调试和验收。")
    add_bullet(doc, "界面要求：以SquareLine Studio生成的LVGL工程为基础，界面标题改为“嘎嘎鸭智能助手”，布局清晰、美观，适合800×480触摸屏。")
    add_bullet(doc, "语音要求：支持点击录音、保存WAV文件、调用百度语音识别，并把识别结果写入问题框。")
    add_bullet(doc, "智能体要求：支持将问题提交到扣子智能体，解析返回内容后在回答区展示。")
    add_bullet(doc, "部署要求：能够通过交叉编译生成ARM可执行文件，并使用SCP/SSH部署到GEC6818开发板运行。")
    add_placeholder(doc, "图1-1 系统任务需求示意图", "可插入课程题目要求截图或最终功能流程图")

    doc.add_heading("1.2 整体设计方案", level=2)
    add_body(doc, "系统总体采用“嵌入式终端 + 云端语音识别 + 云端智能体”的方案。GEC6818开发板负责本地界面、触摸交互、音频采集和结果显示，百度语音识别服务负责把用户语音转换为文本，扣子智能体负责对文本问题进行理解和回答。系统并不在开发板本地部署大语言模型，而是通过HTTP/HTTPS接口调用云端能力，这样可以降低开发板算力和存储压力，同时保留嵌入式设备的人机交互特征。")
    add_body(doc, "软件层面由三部分组成：第一部分是LVGL图形界面层，主要包含按钮、文本框、状态标签和字体加载；第二部分是业务控制层，负责录音线程、发送线程、状态刷新和数据队列；第三部分是接口调用层，包括百度ASR请求、扣子智能体请求、JSON/SSE返回解析和错误处理。主程序循环中持续调用LVGL定时器处理函数和自定义的assistant_ui_poll函数，保证后台线程产生的识别结果和模型回答能够安全同步到界面。")
    add_table(doc, ["模块", "主要功能", "实现要点"], [
        ("UI界面模块", "显示标题、问题、回答和状态", "使用LVGL控件和中文字体，按800×480屏幕布局"),
        ("录音处理模块", "采集用户语音并生成WAV文件", "调用arecord，兼容48kHz双声道采集并转换为16kHz单声道"),
        ("语音识别模块", "将语音文件转换为文本", "通过curl调用百度语音识别接口，解析result字段"),
        ("智能体模块", "把问题发送给大模型并显示回答", "调用扣子v3/chat接口，解析SSE流式answer内容"),
        ("部署调试模块", "编译、上传、运行和定位错误", "使用Make_arm、scp、ssh、日志文件和屏幕提示"),
    ], [2.7, 4.2, 8.0])
    add_placeholder(doc, "图1-2 系统总体结构图", "建议绘制“用户-开发板-LVGL-百度ASR-扣子智能体”的结构图")

    doc.add_heading("1.3 GEC6818平台配置", level=2)
    add_body(doc, "GEC6818开发板作为本系统的运行平台，提供Linux运行环境、触摸屏显示、声卡录音和以太网通信能力。程序运行时通过/dev/fb0访问帧缓冲显示设备，通过/dev/input/event0读取触摸输入，通过arecord调用ALSA录音接口。开发板与电脑处于同一网段，测试中使用的开发板地址为192.168.137.51，网关为电脑共享网络地址192.168.137.1。")
    add_body(doc, "网络配置是云端接口调用能否成功的前提。开发板需要配置IP、子网掩码、默认网关和DNS解析服务。调试时应先确认能ping通电脑网关，再确认能访问外网IP和域名，最后再测试curl访问百度或扣子接口。对于嵌入式课程设计而言，这一步不仅是部署步骤，也是排查“语音识别失败”“智能体请求失败”等问题的重要依据。")
    add_code(doc, """
ifconfig eth0 192.168.137.51 netmask 255.255.255.0 up
route del default
route add default gw 192.168.137.1
echo "nameserver 114.114.114.114" > /etc/resolv.conf
echo "nameserver 223.5.5.5" >> /etc/resolv.conf
ping -c 3 192.168.137.1
ping -c 3 www.baidu.com
""")
    add_placeholder(doc, "图1-3 开发板网络与硬件连接照片", "可插入网线、串口、屏幕和电源连接照片")

    doc.add_heading("2 Linux与嵌入式开发板使用", level=1)
    doc.add_heading("2.1 Linux交叉开发环境配置", level=2)
    add_body(doc, "嵌入式Linux应用通常采用交叉编译方式开发：在PC端编写和编译程序，生成适用于ARM平台的可执行文件，再传输到开发板运行。本工程使用Make_arm作为交叉编译入口，编译器配置为arm-linux-gcc，工程中包含main.c、dict.c、ui目录下的SquareLine Studio导出代码、LVGL源码以及FreeType字体库。")
    add_body(doc, "交叉开发环境配置的关键在于工具链路径、头文件路径和库文件路径。Make_arm中通过CFLAGS加入LVGL目录和freetype_tmp/include/freetype2，通过LDFLAGS链接数学库、FreeType库和pthread库。由于系统采用后台线程分别执行录音和智能体请求，因此pthread链接不可缺少；由于界面需要显示中文，FreeType字体支持也十分重要。")
    add_code(doc, """
cd /mnt/c/嵌入式实操课/lv_port_linux_sdl_gec6818
make --file=Make_arm clean
make --file=Make_arm
""")
    add_body(doc, "编译完成后，目标文件位于build/bin/main。若修改了界面、语音识别或智能体代码，应先清理旧目标文件再重新编译，以避免旧对象文件未更新导致程序行为与源码不一致。调试过程中可以在程序中加入BUILD_MARK字符串，并在开发板上使用strings /main | grep REC_FIX确认当前运行的确实是最新二进制文件。")
    add_table(doc, ["配置项", "配置内容", "作用"], [
        ("交叉编译器", "arm-linux-gcc", "生成开发板可运行的ARM程序"),
        ("图形库", "LVGL", "提供按钮、文本框、标签、屏幕等控件"),
        ("字体库", "FreeType + simkai.ttf", "支持中文显示，避免乱码或方块字"),
        ("线程库", "pthread", "保证录音、网络请求不阻塞主界面"),
        ("网络工具", "curl、scp、ssh", "完成云端接口访问与部署调试"),
    ], [3.0, 4.1, 7.8])

    doc.add_heading("2.2 SSH工具配置", level=2)
    add_body(doc, "开发板调试需要稳定的远程登录方式。PC端通过SSH连接开发板root账户，可以执行网络配置、查看文件、运行程序和查看日志。编译生成的新程序通过scp上传到开发板根目录，若目标文件正在运行，直接覆盖/main可能出现Text file busy错误，因此推荐先上传为/main_new，再在开发板上停止旧进程并重命名。")
    add_code(doc, """
scp ./build/bin/main root@192.168.137.51:/main_new
ssh root@192.168.137.51
cd /
killall main 2>/dev/null
mv /main_new /main
chmod +x /main
./main
""")
    add_body(doc, "在调试过程中，SSH不仅用于运行程序，还用于验证录音设备和接口返回。例如通过which arecord确认录音工具是否存在，通过arecord -l查看捕获设备，通过cat /proc/asound/cards确认声卡编号。实际测试中，开发板的捕获设备为card 0 device 0，但直接使用16kHz单声道可能出现Channels count not available，因此程序增加了48kHz双声道录音与软件转换逻辑。")
    add_placeholder(doc, "图2-1 SSH连接与程序部署截图", "可插入终端中scp、ssh和./main运行截图")

    doc.add_heading("3 屏幕软件设计", level=1)
    doc.add_heading("3.1 主程序设计流程", level=2)
    add_body(doc, "主程序main.c负责初始化LVGL、显示设备、触摸设备、UI界面和中文字体。显示设备通过lv_linux_fbdev_create创建，并绑定到/dev/fb0；触摸输入通过lv_evdev_create创建，并绑定到/dev/input/event0。界面初始化完成后，程序调用apply_chinese_font加载simkai.ttf，使标题、按钮和文本框都能够显示中文。")
    add_body(doc, "主循环采用轻量轮询方式运行，每次循环先调用assistant_ui_poll检查后台线程是否有新的状态、问题或回答需要刷新，再调用lv_timer_handler处理LVGL内部定时任务，最后短暂休眠。这样既保证界面响应，又避免网络请求或录音过程直接阻塞UI线程。整体流程体现了嵌入式图形程序常用的事件驱动思想。")
    add_code(doc, """
while(1) {
    assistant_ui_poll();
    lv_timer_handler();
    usleep(5000);
}
""")
    add_placeholder(doc, "图3-1 主程序运行流程图", "建议绘制初始化、事件循环、后台线程、界面刷新之间的关系")

    doc.add_heading("3.2 UI界面交互程序设计", level=2)
    add_body(doc, "界面由SquareLine Studio设计并导出，再结合课程要求进行二次修改。屏幕顶部显示“嘎嘎鸭智能助手”标题，中间设置模型回答区，下方从左到右依次为“语音输入”按钮、问题输入框和“发送按钮”，底部为状态提示标签。回答区采用较大的多行TextArea，便于展示模型返回的较长内容；问题框采用单行TextArea，便于展示百度语音识别后的问题文本。")
    add_body(doc, "按钮事件分别绑定到RecordAudio和SendQuestion函数。RecordAudio函数用于开始录音，它先清空问题框和回答区，再创建录音线程；SendQuestion函数用于发送问题，它先检查问题是否为空，再复制问题文本并创建智能体请求线程。通过try_begin_job和finish_job维护忙闲状态，避免用户连续点击按钮造成多个录音或网络请求同时运行。")
    add_table(doc, ["控件", "名称", "功能说明"], [
        ("标题标签", "ui_TitleLabel", "显示“嘎嘎鸭智能助手”"),
        ("回答文本框", "ui_TextArea2", "显示录音提示、识别结果提示和智能体回答"),
        ("语音按钮", "ui_Button1", "触发RecordAudio录音流程"),
        ("问题文本框", "ui_TextArea1", "显示识别文字，也可手动输入问题"),
        ("发送按钮", "ui_Button2", "触发SendQuestion智能体请求流程"),
        ("状态标签", "ui_Label1", "显示就绪、录音、识别、请求、失败等状态"),
    ], [2.5, 3.5, 8.0])
    add_placeholder(doc, "图3-2 SquareLine Studio界面设计截图", "可插入SquareLine Studio中的UI设计截图")

    doc.add_heading("3.3 数据显示程序设计", level=2)
    add_body(doc, "由于录音、语音识别和智能体请求都可能耗时数秒，如果直接在LVGL事件回调中执行，会造成界面卡顿甚至无法刷新。因此本设计采用后台线程处理耗时任务，线程完成后通过queue_status、queue_question和queue_answer把待显示内容写入共享缓冲区，并设置dirty标志。主循环中的assistant_ui_poll在UI线程中读取这些标志，再调用lv_label_set_text和lv_textarea_set_text更新界面。")
    add_body(doc, "共享状态使用pthread_mutex_t互斥锁保护，避免后台线程和主线程同时读写同一块内存。界面更新只发生在主线程中，这符合LVGL对象操作的安全要求。问题框更新后还会设置光标位置并滚动到起始处，保证识别文字能够及时显示；回答区则负责展示识别提示、错误原因或智能体回答。")
    add_body(doc, "程序对错误信息也进行了可视化处理。例如录音失败时，会读取/tmp/assistant_arecord.log中的arecord返回内容，并提示用户执行测试命令；语音识别失败时，会把百度接口原始响应写入/tmp/baidu_asr_response.txt；智能体请求失败时，会把扣子原始返回写入/tmp/coze_response.txt。这些日志文件既方便现场调试，也能作为课程验收时说明问题定位过程的材料。")
    add_placeholder(doc, "图3-3 开发板运行界面截图", "可插入开发板屏幕上嘎嘎鸭智能助手运行效果图")

    doc.add_heading("4 大模型部署与应用", level=1)
    doc.add_heading("4.1 大模型部署", level=2)
    add_body(doc, "本设计中的“大模型部署”采用云端智能体部署方式，即在扣子平台创建智能体，配置模型能力、知识库或工作流，并通过开放API供嵌入式程序调用。开发板端只保存智能体ID、用户标识和访问令牌等必要配置，不在本地运行大模型权重文件。这样可以充分利用云端模型的语义理解和回答生成能力，同时使嵌入式端保持较低资源占用。")
    add_body(doc, "为了保证接口调用可控，程序使用curl向https://api.coze.cn/v3/chat发送POST请求，请求体中包含bot_id、user_id、stream和additional_messages等字段。其中additional_messages保存用户问题，stream设置为true表示使用流式响应。实际项目中访问令牌属于敏感信息，论文和截图中不应直接展示完整Token，应在代码中以配置项形式管理，并在提交材料时进行遮挡或删除。")
    add_code(doc, """
{
  "bot_id": "智能体ID",
  "user_id": "gec6818_user",
  "stream": true,
  "additional_messages": [
    {"role": "user", "type": "question", "content_type": "text", "content": "用户问题"}
  ]
}
""")
    add_placeholder(doc, "图4-1 扣子智能体配置界面截图", "请插入智能体基础配置、知识库或发布配置截图，并遮挡令牌")

    doc.add_heading("4.2 大模型应用", level=2)
    add_body(doc, "智能体应用流程发生在用户点击“发送按钮”之后。程序首先从问题输入框读取文本，若为空则提示用户先语音输入或手动输入；若不为空，则创建后台线程执行网络请求。请求返回后，扣子接口会以SSE流式格式输出多段event和data，真正的回答内容位于conversation.message.delta事件对应data中的content字段。")
    add_body(doc, "调试过程中发现，如果只按普通JSON一次性解析，容易把扣子返回中的调试字段、follow_up字段或完整completed字段误认为回答，甚至会因为只读取前几KB内容而漏掉真正的answer。因此程序最终改为逐行读取/tmp/coze_response.txt，先记录当前event，再在data行中判断type是否为answer，仅拼接conversation.message.delta中的content；如果没有delta，则再用conversation.message.completed中的完整回答兜底。该处理方式提高了流式返回解析的稳定性。")
    add_body(doc, "从用户体验看，大模型应用不是单纯把接口返回打印出来，而是要让嵌入式终端形成完整闭环：用户说出问题，系统识别成文字，用户确认并发送，大模型给出回答，屏幕显示最终结果。这个闭环把传统嵌入式界面的按键、文本框和状态提示，与云端大模型的自然语言能力结合起来，体现了嵌入式AI应用的基本形态。")
    add_table(doc, ["阶段", "输入", "输出", "关键处理"], [
        ("问题读取", "TextArea1文本", "question字符串", "检查空文本，复制到线程参数"),
        ("请求构造", "question字符串", "JSON请求体", "转义特殊字符，写入临时请求文件"),
        ("接口调用", "JSON请求体", "SSE响应日志", "curl -k -sS -N POST到扣子接口"),
        ("响应解析", "event/data流", "answer字符串", "只拼接type=answer的delta内容"),
        ("界面刷新", "answer字符串", "TextArea2显示", "通过队列和dirty标志回到UI线程更新"),
    ], [2.2, 3.5, 3.5, 5.0])
    add_placeholder(doc, "图4-2 智能体回答效果截图", "可插入开发板上显示模型回答的照片")

    doc.add_heading("4.3 百度语音识别模型开发", level=2)
    add_body(doc, "语音识别模块用于把用户的语音问题转换为文字，是本系统实现自然交互的关键。程序使用arecord录音，优先尝试生成16kHz、16位、小端、单声道WAV文件；当开发板声卡不支持该参数时，程序改为采集48kHz双声道音频，再在C语言中读取WAV头和PCM数据，将左右声道混合并按采样率比例降采样为16kHz单声道。转换后的文件保存为/tmp/assistant_question.wav。")
    add_body(doc, "百度语音识别接口调用地址为http://vop.baidu.com/pro_api，请求头设置Content-Type: audio/wav;rate=16000，并通过URL参数传入dev_pid、cuid和token。接口返回JSON后，程序解析result数组中的文字内容；若没有结果，则读取err_no和err_msg生成错误提示，并把原始响应保存到/tmp/baidu_asr_response.txt。")
    add_code(doc, """
arecord -D plughw:0,0 -d 5 -f S16_LE -r 16000 -c 1 /tmp/assistant_question.wav
arecord -D plughw:0,0 -d 5 -f S16_LE -r 48000 -c 2 /tmp/assistant_question_stereo.wav
curl -X POST -H "Content-Type: audio/wav;rate=16000" 百度语音识别接口 --data-binary "@/tmp/assistant_question.wav"
""")
    add_body(doc, "在开发板调试中，音频参数适配是最容易出错的环节。通过arecord -l和cat /proc/asound/cards确认设备后，还需要实际录音测试文件大小和格式。若录音文件存在但识别失败，应进一步检查网络、百度Token有效期、音频采样率、WAV头格式和语音内容清晰度。本设计通过日志文件、屏幕提示和备用录音路径，提高了语音识别模块在真实硬件环境下的可调试性。")
    add_placeholder(doc, "图4-3 百度语音识别接口配置与测试截图", "可插入百度控制台、录音文件或识别返回截图，注意遮挡密钥")

    doc.add_heading("5 成果验收与总结", level=1)
    doc.add_heading("5.1 成果验收", level=2)
    add_body(doc, "系统完成后，按照“编译部署—网络测试—录音测试—语音识别—智能体回答—界面展示”的顺序进行验收。首先在PC端使用Make_arm生成main程序，通过scp上传到开发板；其次在开发板上确认网络能访问网关和外部接口；再次点击语音输入按钮，观察是否生成/tmp/assistant_question.wav并返回识别文字；最后点击发送按钮，观察回答区是否显示扣子智能体返回内容。")
    add_table(doc, ["验收项", "预期结果", "验收情况"], [
        ("程序启动", "开发板屏幕显示嘎嘎鸭智能助手界面", "通过"),
        ("触摸交互", "语音输入和发送按钮可响应", "通过"),
        ("录音功能", "生成有效WAV文件并可用于识别", "通过"),
        ("语音识别", "问题框显示百度识别出的中文文字", "通过"),
        ("智能体请求", "回答区显示扣子智能体返回内容", "通过"),
        ("异常提示", "网络、录音或接口异常时能显示提示", "通过"),
    ], [3.0, 7.0, 4.5])
    add_body(doc, "根据调试记录，项目中曾出现录音参数不兼容、网关不通、令牌错误、接口返回解析失败和程序覆盖Text file busy等问题。通过逐项排查，最终形成了较完整的解决方案：录音部分增加48kHz双声道兼容与软件转换；网络部分明确IP、网关和DNS配置；智能体部分更换有效令牌并修正SSE解析；部署部分采用/main_new临时文件替换方式。上述过程说明嵌入式AI应用不仅需要完成代码编写，更需要具备硬件、系统、网络和接口联调能力。")
    add_placeholder(doc, "图5-1 最终成果验收截图", "可插入开发板显示识别文字和模型回答的照片")

    doc.add_heading("5.2 设计总结", level=2)
    add_body(doc, "本次课程设计完成了一个运行在GEC6818开发板上的中文智能语音助手系统。通过LVGL实现图形界面，通过ALSA/arecord完成音频采集，通过百度语音识别完成语音转文字，通过扣子智能体完成问题回答，最终形成了“语音输入—文字识别—智能体思考—屏幕显示”的完整应用链路。该系统虽然规模不大，但覆盖了嵌入式Linux开发中常见的显示、触摸、音频、网络、线程和接口解析等关键内容。")
    add_body(doc, "从技术收获看，本设计加深了对交叉编译、SSH部署、LVGL控件布局、中文字体加载、WAV音频格式、HTTP接口调用和JSON/SSE解析的理解。尤其是在真实开发板调试中，很多问题不是语法错误，而是硬件参数、网络配置或返回格式与预期不一致造成的。通过查看日志、拆分测试命令和逐步验证，可以更快定位问题。")
    add_body(doc, "从改进方向看，后续可继续完善三个方面：一是增加录音倒计时和音量提示，使用户更清楚当前录音状态；二是把百度Token、扣子Token和智能体ID改为外部配置文件，避免重新编译程序；三是加入历史对话列表、语音播报和离线缓存，提高系统实用性。总体而言，本项目实现了课程要求的主要功能，也展示了嵌入式终端接入大模型服务的基本方法。")

    doc.add_heading("参考资料", level=1)
    refs = [
        "【有道云笔记】001-开发环境搭建：https://share.note.youdao.com/s/6uNT0FzT",
        "【有道云笔记】002-LVGL图形界面库开发：https://share.note.youdao.com/s/a93ZN8xL",
        "【有道云笔记】003-百度语音识别大模型：https://share.note.youdao.com/s/Bj55U1OJ",
        "【有道云笔记】004-大模型开发：https://share.note.youdao.com/s/azMDdH0",
        "LVGL官方文档与课程工程源码。",
        "百度智能云语音识别接口说明与扣子开放平台接口说明。",
    ]
    for idx, ref in enumerate(refs, start=1):
        add_noindent(doc, f"[{idx}] {ref}")

    doc.core_properties.title = "嘎嘎鸭智能助手-嵌入式技术课程设计说明书"
    doc.core_properties.subject = "嵌入式大模型应用与开发"
    doc.core_properties.author = "学生"
    doc.core_properties.comments = "由课程设计模板扩写，图片位置已预留。"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
